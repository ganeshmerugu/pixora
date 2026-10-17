from __future__ import annotations
import csv
import uuid
import io
from openpyxl import load_workbook
from pixora.parsing.reader import BaseReader
from pixora.models.document import DocumentTree, Page, Region


class SpreadsheetReader(BaseReader):
    supported_formats = [".xlsx", ".xls"]

    async def read(self, file_bytes: bytes, filename: str) -> DocumentTree:
        pages: list[Page] = []
        wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
        for sheet_idx, sheet_name in enumerate(wb.sheetnames):
            ws = wb[sheet_name]
            rows = []
            cells = []
            for row_idx, row in enumerate(ws.iter_rows(min_row=1), start=1):
                row_values = []
                for cell in row:
                    val = str(cell.value) if cell.value is not None else ""
                    row_values.append(val)
                    cell_meta = {
                        "row": row_idx,
                        "col": cell.column,
                        "col_letter": cell.column_letter,
                        "value": val,
                        "bold": bool(cell.font and cell.font.bold),
                        "font_color": _color_to_rgb(cell.font.color) if cell.font and cell.font.color else None,
                        "bg_color": _color_to_rgb(cell.fill.fgColor) if cell.fill and cell.fill.fgColor else None,
                    }
                    cells.append(cell_meta)
                rows.append(" | ".join(row_values))
            content = "\n".join(rows)
            pages.append(Page(
                page_number=sheet_idx + 1,
                regions=[Region(
                    id=f"r_sheet_{sheet_idx}_{uuid.uuid4().hex[:8]}",
                    type="TABLE",
                    content=content,
                    confidence=1.0,
                    metadata={
                        "sheet_name": sheet_name,
                        "cells": cells,
                        "row_count": ws.max_row or 0,
                        "col_count": ws.max_column or 0,
                    },
                )],
            ))
        wb.close()
        return DocumentTree(
            document_id=uuid.uuid4().hex,
            pages=pages,
            source_format="spreadsheet",
            text_layer_present=True,
        )


def _color_to_rgb(color) -> str | None:
    if color is None:
        return None
    try:
        rgb = color.rgb
        if rgb and isinstance(rgb, str) and len(rgb) >= 6:
            rgb = rgb.lstrip("0")[-6:] if len(rgb) > 6 else rgb
            return f"#{rgb}"
    except Exception:
        pass
    return None


class CSVReader(BaseReader):
    supported_formats = [".csv", ".tsv"]

    async def read(self, file_bytes: bytes, filename: str) -> DocumentTree:
        text = file_bytes.decode("utf-8", errors="replace")
        delimiter = "\t" if filename.endswith(".tsv") else ","
        reader_obj = csv.reader(io.StringIO(text), delimiter=delimiter)
        all_rows = list(reader_obj)
        if not all_rows:
            return DocumentTree(
                document_id=uuid.uuid4().hex,
                pages=[Page(page_number=1, regions=[])],
                source_format="csv",
                text_layer_present=True,
            )

        headers = all_rows[0]
        data_rows = all_rows[1:]

        regions: list[Region] = []

        # Region 1: Key-value summary for each row so dynamic_extract can pick them up
        kv_lines = []
        kv_lines.append(f"total_rows: {len(data_rows)}")
        kv_lines.append(f"columns: {', '.join(headers)}")
        for row_idx, row in enumerate(data_rows):
            for col_idx, hdr in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else ""
                if val.strip():
                    kv_lines.append(f"row_{row_idx + 1}_{_header_to_field(hdr)}: {val}")
        regions.append(Region(
            id=f"r_csv_kv_{uuid.uuid4().hex[:8]}",
            type="TEXT",
            content="\n".join(kv_lines),
            confidence=1.0,
        ))

        # Region 2: Pipe table for table detection
        pipe_lines = [" | ".join(headers)]
        for row in data_rows:
            pipe_lines.append(" | ".join(row))
        regions.append(Region(
            id=f"r_csv_table_{uuid.uuid4().hex[:8]}",
            type="TABLE",
            content="\n".join(pipe_lines),
            confidence=1.0,
            metadata={"headers": headers, "row_count": len(data_rows)},
        ))

        return DocumentTree(
            document_id=uuid.uuid4().hex,
            pages=[Page(page_number=1, regions=regions)],
            source_format="csv",
            text_layer_present=True,
        )


def _header_to_field(header: str) -> str:
    import re
    return re.sub(r"[^a-z0-9]+", "_", header.strip().lower()).strip("_") or "col"
